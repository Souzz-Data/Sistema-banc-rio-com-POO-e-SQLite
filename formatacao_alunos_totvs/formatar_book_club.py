import pandas as pd
import numpy as np
import re
import os
import sys
import warnings
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

# ==========================================
# 1. DEFINIÇÃO DE CAMINHOS FIXOS
# ==========================================
# Uso: python formatar_book_club.py [planilha_entrada.xlsx] [pasta_saida]
arquivo_entrada = r"C:\Users\everton.araujo\Desktop\Disparo_Limpeza\Base_Ativos\book_club.xlsx"
caminho_output = r"C:\Users\everton.araujo\Desktop\Disparo_Limpeza"

if len(sys.argv) > 1: arquivo_entrada = sys.argv[1]
if len(sys.argv) > 2: caminho_output = sys.argv[2]

if not os.path.exists(caminho_output):
    os.makedirs(caminho_output)

NOME_CAMPANHA = 'BookClub_Novembro'

# Grupos de contato: cada um tem nome, e-mail e telefones
CONTATOS = {
    'Resp. Fin.': {
        'nome': 'NOME RESP FIN',
        'email': 'EMAIL RESP FIN',
        'telefones': ['TELEFONE 1 RESP. FIN.', 'TELEFONE 2 RESP. FIN.', 'TELEFONE COMERCIAL RESP. FIN.'],
    },
    'Aluno': {
        'nome': 'NOME',
        'email': 'EMAIL',
        'telefones': ['TELEFONE 1 ALUNO', 'TELEFONE 2 ALUNO', 'TELEFONE 3 ALUNO'],
    },
    'Resp. Acad.': {
        'nome': 'NOME RESP ACA',
        'email': 'EMAIL RESP ACA',
        'telefones': ['TELEFONE 1 RESP. ACAD.', 'TELEFONE 2 RESP. ACAD.', 'TELEFONE 3 RESP. ACAD.'],
    },
}

# Ordem de busca do contato (só celular; fixo nunca entra)
ORDEM_MENOR = ['Resp. Fin.', 'Resp. Acad.']
ORDEM_ADULTO = ['Resp. Fin.', 'Aluno', 'Resp. Acad.']
# Adulto cujo pagante é empresa: usa o contato do próprio aluno (nunca o da empresa)
ORDEM_ADULTO_EMPRESA = ['Aluno', 'Resp. Acad.']

# Pagante pessoa jurídica (empresa/escola que paga o curso)
PADRAO_EMPRESA = r'\bLTDA\b|\bS\.?\s?A\.?$|S/A|EIRELI|ESCOLA|CORRETORA|COMPANY|EMPRESA|CONSULTING|CONSULTORIA'

# Campanha é para alunos acima de 12 anos
IDADE_MINIMA = 13

# Nome da filial no TOTVS -> nome usado nas Tags
RENOMEAR_FILIAIS = {
    'INSTITUTO BRASIL ESTADOS UNIDOS': 'IBEU AT HOME',
}

# Segmento a partir da coluna CURSO (primeira regra que bater)
SEGMENTOS = [
    ('Masterflex', r'MASTER\s*FLEX'),
    ('Top Flex', r'TOP\s*FLEX'),
    ('Conversation', r'CONVERSATION'),
    ('Intermediate', r'INTERMEDIATE'),
    ('Advanced', r'ADVANCED'),
]
ORDEM_SEGMENTOS = ['Intermediate', 'Advanced', 'Masterflex', 'Top Flex', 'Conversation', 'Outros']

VALORES_VAZIOS = ['', 'NAN', 'NONE', 'SEM RESP', 'SEM EMAIL', '(VAZIAS)']

# ==========================================
# 2. FUNÇÕES DE LIMPEZA (TELEFONE, NOME E E-MAIL)
# ==========================================
def formatar_telefone(telefone):
    if pd.isna(telefone) or str(telefone).strip() in ['', '(Vazias)', 'nan']:
        return ''
    telefone_str = str(telefone).strip()
    if telefone_str.startswith('+') and not telefone_str.startswith('+55'):
        return ''
    numeros = re.sub(r'\D', '', telefone_str)
    if not numeros: return ''
    if numeros.startswith('0'): numeros = numeros[1:]
    if numeros.startswith('55') and len(numeros) in [12, 13]: pass
    elif len(numeros) in [10, 11]: numeros = '55' + numeros
    elif len(numeros) in [8, 9]: numeros = '5521' + numeros
    else: return ''
    return numeros

def limpar_nome(nome):
    if pd.isna(nome) or str(nome).strip() == '':
        return ''
    nome_str = str(nome).strip().upper()
    if '@' in nome_str and not nome_str.startswith('@'):
        nome_str = nome_str.split('@')[0].replace('.', ' ').replace('_', ' ')
    nome_str = re.sub(r'^LEAD\s*-\s*', '', nome_str)
    nome_str = re.sub(r'\s*\(\d+\)\s*', '', nome_str)
    nome_str = re.sub(r'[^A-ZÁÀÂÃÉÈÊÍÏÓÔÕÖÚÇÑ\s]', '', nome_str)
    return re.sub(r'\s+', ' ', nome_str).strip()

def is_nome_valido(nome):
    if not nome: return False
    nome_str = str(nome).lower()
    if '{' in nome_str or '}' in nome_str or 'http' in nome_str or 'type:' in nome_str: return False
    if nome_str.startswith('@') or nome_str.startswith('_'): return False
    if re.search(r'\d{4,}', nome_str) or re.search(r'^\(?\d{2}\)?\s*\d+', nome_str): return False
    return True

def limpar_email(email):
    if pd.isna(email):
        return ''
    # Pega o pedaço que contém '@' (ex.: "p fulano@gmail.com" -> "fulano@gmail.com")
    for parte in str(email).strip().split():
        if re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', parte):
            return parte.lower()
    return ''

def is_celular(numero):
    # 55 + DDD + 9 dígitos começando com 9
    return len(numero) == 13 and numero[4] == '9'

def idade_maxima(faixa):
    # "06 - 12 ANOS" -> 12 | "05 ou - ANOS" -> 5 | "61 ou + ANOS" -> sem limite
    faixa_str = str(faixa).upper()
    numeros = [int(n) for n in re.findall(r'\d+', faixa_str)]
    if not numeros or 'OU +' in faixa_str: return 999
    return max(numeros)

def is_menor(faixa):
    # "13 - 17 ANOS" -> menor | "18 - 25 ANOS" / "61 ou + ANOS" -> adulto | "05 ou - ANOS" -> menor
    faixa_str = str(faixa).upper()
    numeros = [int(n) for n in re.findall(r'\d+', faixa_str)]
    if not numeros: return False
    if 'OU +' in faixa_str: return False
    return max(numeros) < 18

def definir_segmento(curso):
    curso_str = str(curso).upper()
    for segmento, padrao in SEGMENTOS:
        if re.search(padrao, curso_str):
            return segmento
    return 'Outros'

def buscar_contato(linha, ordem):
    """Busca celular e e-mail separadamente, seguindo a ordem dos grupos."""
    celular, origem_cel, email, origem_email = '', '', '', ''
    for grupo in ordem:
        cfg = CONTATOS[grupo]
        if not celular:
            for col in cfg['telefones']:
                tel = formatar_telefone(linha.get(col))
                if tel and is_celular(tel):
                    celular, origem_cel = tel, grupo
                    break
        if not email:
            email = limpar_email(linha.get(cfg['email']))
            if email: origem_email = grupo
    return celular, origem_cel, email, origem_email

def nome_do_grupo(linha, grupo):
    nome = limpar_nome(linha.get(CONTATOS[grupo]['nome']))
    return nome if nome.upper() not in VALORES_VAZIOS and is_nome_valido(nome) else ''

# ==========================================
# 3. LEITURA DA BASE
# ==========================================
print("Lendo base do Book Club (TOTVS)...")
try:
    df = pd.read_excel(arquivo_entrada)
except FileNotFoundError as e:
    print(f"❌ Erro ao ler a base: {e}")
    exit()

df.columns = [str(c).strip() for c in df.columns]
qtd_original = len(df)

df = df[df['FAIXA ETÁRIA'].apply(idade_maxima) >= IDADE_MINIMA]
qtd_abaixo_idade = qtd_original - len(df)

# ==========================================
# 4. MONTAGEM DOS DADOS POR ALUNO
# ==========================================
registros = []
qtd_empresa = 0
for _, linha in df.iterrows():
    menor = is_menor(linha.get('FAIXA ETÁRIA'))
    pagante_empresa = bool(re.search(PADRAO_EMPRESA, str(linha.get('NOME RESP FIN', '')).upper()))
    if menor: ordem = ORDEM_MENOR
    elif pagante_empresa: ordem = ORDEM_ADULTO_EMPRESA
    else: ordem = ORDEM_ADULTO
    celular, origem_cel, email, origem_email = buscar_contato(linha, ordem)
    qtd_empresa += int(pagante_empresa and not menor)

    # Nome do pagante; se não houver, usa o nome de quem forneceu o contato
    responsavel = nome_do_grupo(linha, 'Resp. Fin.')
    if not responsavel and (origem_cel or origem_email):
        responsavel = nome_do_grupo(linha, origem_cel or origem_email)

    nome_aluno = str(linha.get('NOME', '')).strip()
    ra = re.search(r'\((\d+)\)', nome_aluno)
    registros.append({
        'Filial': RENOMEAR_FILIAIS.get(str(linha.get('FILIAL', '')).strip(), str(linha.get('FILIAL', '')).strip()),
        'Segmento': definir_segmento(linha.get('CURSO')),
        'Curso': str(linha.get('CURSO', '')).strip(),
        'Nível': str(linha.get('NIVEL', '')).strip(),
        'RA': ra.group(1) if ra else '',
        'Aluno': limpar_nome(nome_aluno),
        'Faixa Etária': str(linha.get('FAIXA ETÁRIA', '')).strip(),
        'Idade': 'Menor' if menor else 'Adulto',
        'Responsável Pagante': responsavel,
        'E-mail': email,
        'Celular': celular,
        'Origem do E-mail': origem_email,
        'Origem do Celular': origem_cel,
    })

df_detalhe = pd.DataFrame(registros)
df_detalhe['Segmento'] = pd.Categorical(df_detalhe['Segmento'], ORDEM_SEGMENTOS, ordered=True)
df_detalhe = df_detalhe.sort_values(['Filial', 'Segmento', 'Aluno']).reset_index(drop=True)
df_detalhe['Segmento'] = df_detalhe['Segmento'].astype(str)

# ==========================================
# 5. BASE DE DISPARO (MOLDE FILIAIS)
# ==========================================
df_disparo = df_detalhe[df_detalhe['Celular'] != ''].copy()
qtd_com_celular = len(df_disparo)

df_molde = pd.DataFrame()
df_molde['Nome'] = df_disparo['Aluno']
df_molde['Telefone'] = df_disparo['Celular']
df_molde['country code'] = '55'
df_molde['Tags'] = df_disparo['Filial']
df_molde['E-mail'] = df_disparo['E-mail'].replace('', np.nan)

# Mesmo pagante com mais de um filho recebe uma vez só
df_molde = df_molde.drop_duplicates(subset=['Telefone'], keep='first')
qtd_duplicados = qtd_com_celular - len(df_molde)

# ==========================================
# 6. EXPORTAÇÃO
# ==========================================
COR_CABECALHO = '1F4E78'
COR_ZEBRA = 'EAF1F8'
COR_SEM_CONTATO = 'FCE4D6'
borda = Border(bottom=Side(style='thin', color='BFBFBF'))

def formatar_aba(ws, df_aba, destacar_sem_celular=False):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cel in ws[1]:
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor=COR_CABECALHO)
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30

    col_celular = list(df_aba.columns).index('Celular') + 1 if 'Celular' in df_aba.columns else None
    for i, linha in enumerate(ws.iter_rows(min_row=2), start=0):
        sem_celular = destacar_sem_celular and col_celular and not linha[col_celular - 1].value
        cor = COR_SEM_CONTATO if sem_celular else (COR_ZEBRA if i % 2 else None)
        for cel in linha:
            cel.border = borda
            cel.alignment = Alignment(vertical='center')
            if cor: cel.fill = PatternFill('solid', fgColor=cor)

    # Telefones como texto e largura de colunas ajustada ao conteúdo
    for idx, coluna in enumerate(df_aba.columns, start=1):
        letra = get_column_letter(idx)
        if coluna in ['Celular', 'Telefone', 'RA', 'country code']:
            for cel in ws[letra][1:]:
                cel.number_format = '@'
        maior = max([len(str(coluna))] + [len(str(v)) for v in df_aba[coluna].fillna('')])
        ws.column_dimensions[letra].width = min(max(maior + 3, 10), 50)

def escrever_aba(writer, nome_aba, df_aba, destacar_sem_celular=False):
    df_aba.to_excel(writer, sheet_name=nome_aba, index=False)
    formatar_aba(writer.sheets[nome_aba], df_aba, destacar_sem_celular)

# Planilha detalhada: aba Todos + uma aba por segmento
arquivo_detalhado = os.path.join(caminho_output, f"{NOME_CAMPANHA}_Detalhada.xlsx")
with pd.ExcelWriter(arquivo_detalhado, engine='openpyxl') as writer:
    escrever_aba(writer, 'Todos', df_detalhe, destacar_sem_celular=True)
    for segmento in ORDEM_SEGMENTOS:
        df_seg = df_detalhe[df_detalhe['Segmento'] == segmento]
        if len(df_seg):
            escrever_aba(writer, segmento, df_seg, destacar_sem_celular=True)

# Base de disparo
arquivo_disparo = os.path.join(caminho_output, f"{NOME_CAMPANHA}_Disparo.xlsx")
with pd.ExcelWriter(arquivo_disparo, engine='openpyxl') as writer:
    escrever_aba(writer, 'Disparo', df_molde)

# ==========================================
# 7. RELATÓRIO
# ==========================================
sem_celular = (df_detalhe['Celular'] == '').sum()
sem_email = (df_detalhe['E-mail'] == '').sum()
sem_nada = ((df_detalhe['Celular'] == '') & (df_detalhe['E-mail'] == '')).sum()

print("\n" + "="*60)
print(f"✅ SUCESSO! Book Club formatado.")
print(f"📁 Detalhada: {arquivo_detalhado}")
print(f"📁 Disparo:   {arquivo_disparo}")
print("="*60)

print("\n📊 RELATÓRIO:")
print("-" * 60)
print(f"📉 Total de alunos na base: {qtd_original}")
print(f"🚫 Abaixo de {IDADE_MINIMA} anos (removidos): -{qtd_abaixo_idade}")
print(f"🏢 Adultos com pagante empresa (contato do próprio aluno): {qtd_empresa}")
print(f"🧒 Menores (contato do Resp. Fin./Acad.): {(df_detalhe['Idade'] == 'Menor').sum()}")
print(f"🧑 Adultos (Resp. Fin. → Aluno → Resp. Acad.): {(df_detalhe['Idade'] == 'Adulto').sum()}")
print(f"☎️ Sem celular (só fixo ou sem telefone): {sem_celular}")
print(f"📧 Sem e-mail: {sem_email}")
print(f"🚫 Sem celular e sem e-mail: {sem_nada}")
print(f"👯 Cortados no disparo por mesmo celular (irmãos/mesmo pagante): -{qtd_duplicados}")
print("-" * 60)
print(f"🚀 TOTAL PRONTO PARA DISPARO: {len(df_molde)} contatos")
print("="*60)

print("\n📚 ALUNOS POR SEGMENTO (PLANILHA DETALHADA):")
print("-" * 60)
for segmento, quantidade in df_detalhe['Segmento'].value_counts().reindex(ORDEM_SEGMENTOS).dropna().items():
    print(f"🔹 {segmento}: {int(quantidade)} aluno(s)")

print("\n🏢 CONTATOS NO DISPARO POR FILIAL (TAGS):")
print("-" * 60)
for filial, quantidade in df_molde['Tags'].value_counts().items():
    print(f"🔹 {filial}: {quantidade} contato(s)")
print("-" * 60)
